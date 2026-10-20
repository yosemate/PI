from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re, json

ROOT=Path(__file__).resolve().parent
ASSETS=ROOT/'figures'; ASSETS.mkdir(exist_ok=True)
FONT=Path('C:/Windows/Fonts/arial.ttf')
BOLD=Path('C:/Windows/Fonts/arialbd.ttf')
def font(size=34,bold=False): return ImageFont.truetype(str(BOLD if bold else FONT),size)
def canvas(w,h):
    im=Image.new('RGB',(w,h),'white'); return im,ImageDraw.Draw(im)
def lines(draw,text,maxwidth,size=34,bold=False):
    out=[]
    for para in text.split('\n'):
        current=''
        for word in para.split():
            test=(current+' '+word).strip()
            if draw.textlength(test,font=font(size,bold))>maxwidth and current:
                out.append(current); current=word
            else: current=test
        out.append(current)
    return out
def label(d,box,text,size=34,bold=False):
    x1,y1,x2,y2=box; ts=lines(d,text,x2-x1-30,size,bold); step=size+9
    yy=(y1+y2-len(ts)*step)/2
    for t in ts:
        xx=(x1+x2-d.textlength(t,font=font(size,bold)))/2
        d.text((xx,yy),t,font=font(size,bold),fill='black'); yy+=step
def rect(d,b,text,size=34):
    d.rectangle(b,outline='black',width=3); label(d,b,text,size)
def wire(d,pts,fill='black',width=4): d.line(pts,fill=fill,width=width)
def arrow(d,a,b):
    d.line([a,b],fill='black',width=4)
    x,y=b; dx,dy=b[0]-a[0],b[1]-a[1]
    if abs(dx)>abs(dy):
        s=1 if dx>0 else -1; d.polygon([(x,y),(x-18*s,y-10),(x-18*s,y+10)],fill='black')
    else:
        s=1 if dy>0 else -1; d.polygon([(x,y),(x-10,y-18*s),(x+10,y-18*s)],fill='black')
def save(im,name): im.save(ASSETS/f'{name}.png')

im,d=canvas(1600,820)
steps=[('Material\nAmostra',(60,70,350,220)),('Analisador NIR\nEspectro',(450,70,750,220)),('Aquisição\nDados e identificação',(850,70,1160,220)),('Pré-processamento',(1210,70,1540,220)),('Modelo PLS\nCalibrado e validado',(1210,380,1540,545)),('Kappa estimado\nQualidade e horário',(820,380,1150,545)),('CLP ou SCD/SDCD\nControle',(440,380,750,545)),('SCADA\nSupervisão e histórico',(60,380,350,545))]
for t,b in steps:rect(d,b,t,32)
for a,b in [((350,145),(450,145)),((750,145),(850,145)),((1160,145),(1210,145)),((1375,220),(1375,380)),((1210,460),(1150,460)),((820,460),(750,460)),((440,460),(350,460))]:arrow(d,a,b)
rect(d,(480,650,1150,775),'Referências laboratoriais\nConstrução e verificação do modelo',32)
wire(d,[(1150,710),(1380,710),(1380,545)])
save(im,'industrial')

im,d=canvas(1800,1450)
rect(d,(80,100,570,900),'ESP32\nPlaca com USB tipo C',40)
rect(d,(1180,70,1700,400),'GY-AS7263\nSensor óptico',40)
rect(d,(1180,530,1700,790),'Servo da gaveta\nCurso nominal de 180°',38)
rect(d,(1180,960,1700,1220),'Servo do suporte\nCurso nominal de 180°',38)
rect(d,(80,1060,570,1320),'Fonte externa dos servos\nSamsung EP-TA200B\n5 V / 2 A (etiqueta)',34)
for y,t,target in [(180,'3V3',180),(265,'GPIO21 / SDA',260),(350,'GPIO22 / SCL',340)]:
    label(d,(580,y-45,1100,y+5),t,30)
    wire(d,[(570,y),(1180,target)])
label(d,(590,455,1050,510),'GPIO18 / sinal',30)
wire(d,[(570,520),(820,520),(820,600),(1180,600)])
label(d,(590,610,1050,670),'GPIO19 / sinal',30)
wire(d,[(570,690),(930,690),(930,1020),(1180,1020)])
wire(d,[(570,1125),(1080,1125),(1080,690),(1180,690)],fill='#555555')
wire(d,[(1080,1125),(1180,1125)],fill='#555555')
label(d,(630,1160,1030,1240),'+5 V dos servos',30)
# GND star: the three charger-negative branches meet at one junction.
star=(730,1340)
wire(d,[(325,1320),(325,1340),star])
wire(d,[star,(730,820),(570,820)])
wire(d,[star,(1740,1340),(1740,750),(1700,750)])
wire(d,[star,(1740,1340),(1740,1180),(1700,1180)])
d.ellipse((720,1330,740,1350),fill='black')
wire(d,[(1180,390),(1030,390),(1030,820),(730,820)])
label(d,(70,920,610,1000),'GND comum da placa e do sensor',28)
label(d,(790,1350,1740,1430),'Negativo da fonte: 3 ramais para ESP32 e os dois servos',29)
label(d,(1240,805,1700,885),'Sinal • +5 V • GND',28)
label(d,(1180,1240,1710,1300),'Sinal • +5 V • GND',28)
save(im,'electrical')

im,d=canvas(1600,1000)
d.rectangle((190,130,1380,780),outline='black',width=5)
label(d,(200,30,1400,115),'Estrutura final em MDF com interior preto fosco',36,True)
wire(d,[(370,660),(370,260),(1220,260),(1220,660)],width=9)
rect(d,(690,275,920,400),'Suporte móvel\nReferência: talha',28)
rect(d,(730,460,880,570),'Sensor',30)
wire(d,[(805,400),(805,460)],width=4)
rect(d,(430,650,1150,755),'Gaveta de amostras',34)
rect(d,(710,610,900,650),'Amostra',25)
arrow(d,(805,570),(805,610))
rect(d,(240,340,445,450),'Servo do\nsuporte',27)
rect(d,(1210,650,1460,755),'Servo da\ngaveta',27)
label(d,(920,385,1380,520),'Suporte acima da gaveta\nTransmissão não detalhada',28)
label(d,(210,825,1390,970),'Representação funcional sem escala.\nPosições dos servos são ilustrativas; não se especificam cotas, folgas ou capacidade de carga.',32)
save(im,'physical')

im,d=canvas(1600,900)
rect(d,(70,100,480,320),'Computador\nSupervisório SpectrumLab',34)
rect(d,(620,100,1030,320),'ESP32\nFirmware e sequenciamento',34)
rect(d,(1150,100,1530,320),'GY-AS7263\nSeis canais ópticos',34)
arrow(d,(480,170),(620,170));arrow(d,(620,245),(480,245))
label(d,(440,335,1050,435),'USB serial • Bluetooth • TCP próprio',30)
arrow(d,(1030,180),(1150,180));arrow(d,(1150,250),(1030,250))
rect(d,(560,530,900,710),'Servo da gaveta',34)
rect(d,(1050,530,1450,710),'Servo do suporte',34)
arrow(d,(760,320),(760,530));wire(d,[(960,320),(960,470),(1250,470)]);arrow(d,(1250,470),(1250,530))
label(d,(70,740,1520,870),'O sensor envia dados ao ESP32. O computador processa e apresenta os resultados.\nA bancada não representa uma integração implementada com CLP industrial.',31)
save(im,'prototype')

im,d=canvas(1500,1800)
ys=[80,330,580,830,1080,1330]
texts=['Comunicação e identificação do sensor','Referência óptica válida?\nSe não: executar branco e escuro','Inserir amostra e fechar gaveta','Posicionar conjunto óptico e adquirir canais','Corrigir dados e apresentar resultado demonstrativo','Recolher conjunto e permitir retirada da amostra']
for y,t in zip(ys,texts):rect(d,(220,y,1280,y+175),t,38)
for y in ys[:-1]:arrow(d,(750,y+175),(750,y+250))
label(d,(100,1560,1400,1760),'Falha de leitura ou comunicação: interromper o ciclo e indicar o estado.\nParada comandada por software não equivale a corte elétrico de segurança.',34)
save(im,'flow')

doc=Document()
sec=doc.sections[0]; sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(3); sec.left_margin=Cm(3);sec.bottom_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Heading 4','Caption']:
    s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(12 if name!='Caption' else 10)
for name in ['TOC 1','TOC 2','TOC 3','TOC 4']:
    try:s=doc.styles[name]
    except KeyError:s=doc.styles.add_style(name,1)
    s.font.name='Arial';s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.line_spacing=1.0;s.paragraph_format.space_after=Pt(5)
normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.first_line_indent=Cm(1.25)
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2','Heading 3','Heading 4']:
    s=doc.styles[name];s.font.bold=True;s.paragraph_format.first_line_indent=Cm(0)
    s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(8)
    s.paragraph_format.line_spacing=1.2;s.paragraph_format.keep_with_next=True
    s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['Caption'].paragraph_format.first_line_indent=Cm(0)
doc.styles['Caption'].paragraph_format.line_spacing=1.0
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.core_properties.title='Aplicação da tecnologia NIR na predição do número Kappa'
doc.core_properties.author='Davi Baracat Gonçalves; Kauê Uã Carvalho de Alencar; Mateus Ferreira da Silva'
def p(text='',style=None):return doc.add_paragraph(text,style)
def flat(text='',align=WD_ALIGN_PARAGRAPH.LEFT):
    t=p(text);t.paragraph_format.first_line_indent=Cm(0);t.alignment=align;return t
def centered(text,bold=False,space=0):
    t=flat(text,WD_ALIGN_PARAGRAPH.CENTER);t.paragraph_format.space_before=Pt(space)
    if bold:
        for r in t.runs:r.bold=True
    return t
title='APLICAÇÃO DA TECNOLOGIA NIR NA PREDIÇÃO DO NÚMERO KAPPA'
for t in ['SERVIÇO NACIONAL DE APRENDIZAGEM INDUSTRIAL','SENAI TRÊS LAGOAS','TÉCNICO EM AUTOMAÇÃO INDUSTRIAL']:centered(t,True)
centered('DAVI BARACAT GONÇALVES',space=85)
centered('KAUÊ UÃ CARVALHO DE ALENCAR')
centered('MATEUS FERREIRA DA SILVA')
centered(title,True,95)
centered('TRÊS LAGOAS MS',space=190);centered('2026')
doc.add_page_break()
for t in ['DAVI BARACAT GONÇALVES','KAUÊ UÃ CARVALHO DE ALENCAR','MATEUS FERREIRA DA SILVA']:centered(t)
centered(title,True,110)
t=flat('Trabalho de Conclusão de Curso apresentado ao SENAI Três Lagoas como requisito para obtenção do título de Técnico em Automação Industrial.')
t.paragraph_format.left_indent=Cm(7);t.paragraph_format.space_before=Pt(65);t.paragraph_format.line_spacing=1.0
t=flat('Orientador(a): __________________________');t.paragraph_format.left_indent=Cm(7);t.paragraph_format.space_before=Pt(18)
centered('TRÊS LAGOAS MS',space=190);centered('2026')
def front_heading(text):
    doc.add_page_break(); t=flat(text,WD_ALIGN_PARAGRAPH.CENTER);t.runs[0].bold=True;t.paragraph_format.space_after=Pt(20)
front_heading('RESUMO')
summary='Este trabalho analisa a aplicação da espectroscopia no infravermelho próximo à predição do número Kappa e sua relação com a automação industrial. O estudo apresenta os fundamentos da produção de celulose, da lignina residual, da aquisição espectral, da quimiometria e da regressão por mínimos quadrados parciais. A pesquisa utiliza procedimentos bibliográficos e documentais, com análise de aplicações publicadas e documentação técnica de instrumentação e comunicação. Os resultados examinados mostram que modelos espectrais podem relacionar medições ópticas a valores Kappa de referência em condições definidas de material, preparação e validação. A discussão distingue resultados científicos, funcionalidades comerciais de analisadores e requisitos de utilização da informação no processo. Também apresenta o protótipo didático SpectrumLab, formado por sensor multicanal, ESP32, dois servomotores e supervisório. A documentação inclui esquema elétrico, arranjo físico funcional, arquitetura e sequência operacional. A bancada utiliza seis canais entre 610 e 860 nm e equações fixas para exibir um resultado demonstrativo; não possui modelo PLS treinado com dados próprios nem validação com Kappa laboratorial. Conclui-se que a confiabilidade da predição depende da cadeia completa de medição, calibração e avaliação, enquanto o protótipo contribui para explicar o fluxo de aquisição e supervisão sem substituir um analisador industrial.'
flat(summary).alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
flat('Palavras-chave: espectroscopia NIR; número Kappa; regressão PLS; automação industrial; protótipo didático.').paragraph_format.space_before=Pt(12)
front_heading('ABSTRACT')
abstract='This work examines the application of near-infrared spectroscopy to Kappa number prediction and its relationship with industrial automation. It presents the fundamentals of pulp production, residual lignin, spectral acquisition, chemometrics and partial least squares regression. The study uses bibliographic and documentary procedures to analyze published applications and technical documentation on instrumentation and communication. The reviewed findings show that spectral models can relate optical measurements to reference Kappa values under defined conditions of material, preparation and validation. The discussion distinguishes scientific results, commercial analyzer features and requirements for using analytical information in process operation. It also describes the SpectrumLab educational prototype, consisting of a multichannel sensor, an ESP32, two servomotors and supervisory software. Its documentation includes electrical connections, a functional physical layout, system architecture and an operating sequence. The prototype uses six channels from 610 to 860 nm and fixed equations to display a demonstration result; it has neither a PLS model trained on original data nor validation against laboratory Kappa values. The work concludes that prediction reliability depends on the complete measurement, calibration and evaluation chain, while the prototype illustrates acquisition and supervision without replacing an industrial analyzer.'
flat(abstract).alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
flat('Keywords: NIR spectroscopy; Kappa number; PLS regression; industrial automation; educational prototype.').paragraph_format.space_before=Pt(12)
front_heading('LISTA DE ABREVIATURAS E SIGLAS')
for t in ['CLP — Controlador lógico programável','GND — Referência elétrica comum','GPIO — Entrada ou saída de propósito geral','I²C — Circuito interintegrado, barramento de comunicação','MDF — Painel de fibras de madeira de média densidade','NIR — Infravermelho próximo','OPC UA — Open Platform Communications Unified Architecture','PLS — Mínimos quadrados parciais','PLSR — Regressão por mínimos quadrados parciais','R² — Coeficiente de determinação','RMSECV — Raiz do erro quadrático médio de validação cruzada','RMSEP — Raiz do erro quadrático médio de predição','RPD — Razão entre dispersão de referência e erro preditivo','SCADA — Sistema de supervisão, controle e aquisição de dados','SCD ou SDCD — Sistema de controle distribuído','SNV — Standard Normal Variate','SVM — Máquina de vetores de suporte','TCP — Protocolo de controle de transmissão','VIS — Região visível do espectro']:flat(t)
front_heading('LISTA DE ILUSTRAÇÕES')
listcaptions=['Figura 1 — Fluxo funcional da predição do Kappa e da integração à automação','Figura 2 — Esquema das ligações elétricas do protótipo','Figura 3 — Arranjo físico funcional da bancada e do suporte superior do sensor','Figura 4 — Arquitetura funcional do protótipo SpectrumLab','Figura 5 — Fluxograma funcional de aquisição e supervisão da bancada','Figura 6 — Interface do SpectrumLab em estado de espera e sem aquisição ativa','Quadro 1 — Estudos de predição de Kappa por NIR e PLS','Quadro 2 — Alcance didático e analítico do protótipo','Quadro 3 — Conexões documentadas do protótipo']
for i,t in enumerate(listcaptions):
    par=flat(t+'  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),f'PAGEREF illustration{i+1} \\h');par._p.append(fld)
front_heading('SUMÁRIO')
tocpara=flat(); fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-4" \\h \\z \\u')
rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='Sumário';rr.append(tt);fld.append(rr);tocpara._p.append(fld)
# Numeric pages begin in body, while preliminary pages remain unnumbered.
new=doc.add_section();new.top_margin=Cm(3);new.left_margin=Cm(3);new.bottom_margin=Cm(2);new.right_margin=Cm(2)
new.header.is_linked_to_previous=False
hp=new.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.first_line_indent=Cm(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');hp._p.append(field)
new.footer.is_linked_to_previous=False

def hyperlink(par,text,url):
    rel=par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);par._p.append(h)
def textpara(text,reference=False):
    t=p()
    if reference:
        t.paragraph_format.first_line_indent=Cm(0);t.paragraph_format.line_spacing=1.0;t.paragraph_format.space_after=Pt(12);t.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pos=0
    for m in re.finditer(r'https://[^\s]+',text):
        t.add_run(text[pos:m.start()]);u=m.group().rstrip('.');hyperlink(t,u,u)
        if m.group().endswith('.'):t.add_run('.')
        pos=m.end()
    t.add_run(text[pos:]);return t
def bookmark(par,index):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(index));start.set(qn('w:name'),f'illustration{index}')
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(index));par._p.insert(0,start);par._p.append(end)
def table(title,headers,rows,widths):
    cap=p(title,'Caption');cap.paragraph_format.keep_with_next=True
    bookmark(cap,6+int(re.search(r'Quadro (\d)',title)[1]))
    tab=doc.add_table(rows=1,cols=len(headers));tab.autofit=False
    for i,w in enumerate(widths):tab.columns[i].width=Cm(w)
    for i,h in enumerate(headers):tab.rows[0].cells[i].text=h
    for row in rows:
        cells=tab.add_row().cells
        for c,val in zip(cells,row):c.text=str(val)
    for ri,row in enumerate(tab.rows):
        trpr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trpr.append(keep)
        if ri==0:
            rep=OxmlElement('w:tblHeader');trpr.append(rep)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                elem=OxmlElement('w:'+side);elem.set(qn('w:val'),'single');elem.set(qn('w:sz'),'4');elem.set(qn('w:color'),'D9D9D9');b.append(elem)
            pr.append(b)
            marg=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');marg.append(x)
            pr.append(marg)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAEAEA');pr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.first_line_indent=Cm(0);par.paragraph_format.line_spacing=1.0;par.paragraph_format.space_after=Pt(4);par.alignment=WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.keep_with_next=ri<len(tab.rows)-1
                for r in par.runs:r.font.size=Pt(10);r.bold=ri==0
    p('Fonte: elaboração dos autores.' if 'Quadro 1' not in title else 'Fonte: Fardim, Ferreira e Durán (2002); Zhang et al. (2012); Santos et al. (2014).','Caption')

inrefs=False
for block in (ROOT/'trabalho.md').read_text(encoding='utf-8').split('\n\n'):
    block=block.strip()
    if not block:continue
    if block.startswith('#'):
        for line in block.splitlines():
            m=re.match(r'^(#+) (.*)',line)
            if m:
                depth=len(m[1]);txt=m[2]
                inrefs=txt=='REFERÊNCIAS'
                if txt.startswith('APÊNDICE'):inrefs=False
                doc.add_heading(txt,level=min(depth,4))
            elif line:textpara(line,inrefs)
    elif block.startswith('[[FIG:'):
        name,cap,src=block[6:-2].split('|')
        t=p(cap,'Caption');t.paragraph_format.keep_with_next=True
        bookmark(t,int(re.search(r'Figura (\d)',cap)[1]))
        t=flat('',WD_ALIGN_PARAGRAPH.CENTER);t.paragraph_format.keep_with_next=True
        path=(ASSETS/f'{name}.png') if name!='interface' else Path('C:/Users/mateu/.gemini/antigravity-ide/scratch/spectrum_lab/validation/interface_4_13_0_light_1366.png')
        t.add_run().add_picture(str(path),width=Cm(15.7))
        s=p(src,'Caption');s.paragraph_format.space_after=Pt(10)
    elif block=='[[TABLE:studies]]':
        table('Quadro 1 Estudos de predição de Kappa por NIR e PLS',['Estudo','Material e método','Avaliação documentada'],[
            ['Fardim et al. (2002)','Polpa Kraft de eucalipto; NIR em reflectância e PLS','Modelagem de Kappa e outras propriedades; validação cruzada'],
            ['Zhang et al. (2012)','Polpa e pó de madeira de Acacia spp.; NIR e PLS','54 amostras de calibração e 19 de validação'],
            ['Santos et al. (2014)','Polpa de Acacia melanoxylon; NIR–PLSR','48 amostras de desenvolvimento e 27 de teste; RMSEP 0,4; R² 0,921']], [3.1,6,6.9])
    elif block=='[[TABLE:comparison]]':
        table('Quadro 2 Alcance didático e analítico do protótipo',['Função','Condição documentada','Alcance'],[
            ['Aquisição óptica','Seis canais de 610 a 860 nm','Demonstração de resposta multicanal'],
            ['Referência óptica','Branco e escuro por modo','Correção instrumental relativa'],
            ['Tratamento','Normalização pela soma','Não equivale à transformação SNV'],
            ['Resultado Kappa','Equações fixas demonstrativas','Não validado contra Kappa laboratorial'],
            ['PLS','Discutido na fundamentação','Não treinado com amostras próprias'],
            ['Movimentação','Dois servos; posições estimadas','Sequência automatizada sem fim de curso'],
            ['Supervisão','Software de bancada e comunicação própria','Demonstração; sem integração fabril comprovada']], [3.2,6.4,6.4])
    elif block=='[[TABLE:pins]]':
        table('Quadro 3 Conexões documentadas do protótipo',['Origem','Destino','Função'],[
            ['ESP32 3V3','Alimentação do módulo óptico','Alimentação de 3,3 V'],
            ['ESP32 GPIO21','SDA do módulo óptico','Dados I²C'],
            ['ESP32 GPIO22','SCL do módulo óptico','Relógio I²C'],
            ['ESP32 GPIO18','Sinal do servo da gaveta','Comando de posição'],
            ['ESP32 GPIO19','Sinal do servo do suporte','Comando de posição'],
            ['Fonte externa +5 V','Positivo dos dois servos','Alimentação dos atuadores'],
            ['Fonte externa negativo','GND ESP32 e GND dos dois servos','Três ramais com referência comum'],
            ['GND do módulo óptico','GND da ESP32','Referência comum do sensor']], [5,5.7,5.3])
    else:textpara(block,inrefs)

settings=doc.settings.element;up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');settings.append(up)
out=ROOT/'PI_NIR_Kappa_Rev_0.3.docx';doc.save(out)
print(json.dumps({'output':str(out),'paragraphs':len(doc.paragraphs),'words':len((ROOT/'trabalho.md').read_text(encoding='utf-8').split()),'figures':6,'tables':3,'references':24}))
